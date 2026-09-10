#!/usr/bin/env python3
"""Comprueba el bloque de celda_jupyter.py contra las cifras de control.

Ejecuta la lectura de los dos Excel del cierre de junio 2026 sin abrir ninguna
ventana, y contrasta contra los importes verificados uno por uno.

    python verificar.py Balanza_062026.xlsx ResultadosQES.xlsb
"""

import sys
import tempfile
from pathlib import Path

AQUI = Path(__file__).resolve().parent

# ---------------------------------------------------------------------------
# El bloque se ejecuta tal cual, salvo la última línea que abre la ventana.
# ---------------------------------------------------------------------------
FUENTE = (AQUI / "celda_jupyter.py").read_text(encoding="utf-8")
assert FUENTE.count("\nabrir_ventana()\n") == 1, "cambió el pie del bloque"
BLOQUE = {}
exec(compile(FUENTE.replace("\nabrir_ventana()\n", "\n"), "celda_jupyter.py", "exec"), BLOQUE)

Historico = BLOQUE["Historico"]
construir_html = BLOQUE["construir_html"]
escribir_vista = BLOQUE["escribir_vista"]
etiqueta_corta = BLOQUE["etiqueta_corta"]

# ---------------------------------------------------------------------------
# Cifras de control (sección 05 del encargo), en USD
# ---------------------------------------------------------------------------
ESPERADO = {
    #  periodo        rrc local      rrc cnsf       rsr (los dos)  rsnr local  rsnr cnsf
    "2025-07-31": (8_331_317.86, 11_142_522.97, 2_818_658.21, 12_701.88, 325_011.19),
    "2025-09-30": (8_723_257.33, 11_817_782.40, 3_722_728.06, 12_701.88, 423_226.93),
    "2025-11-30": (9_132_774.22, 12_231_493.45, 4_246_854.22, 12_701.88, 599_138.64),
    "2025-12-31": (9_278_886.43, 12_604_139.42, 3_890_396.27,  7_461.06, 419_827.70),
    "2026-03-31": (9_171_744.38, 12_718_231.37, 4_187_990.32,  7_461.06, 555_998.70),
    "2026-06-30": (9_084_601.70, 12_799_633.26, 2_773_742.22,  7_461.06, 809_405.38),
}

# La vista, en millones de USD: (local, cnsf, dif) por corte y el incremento
VISTA = {
    "2025-12-31": {"rrc": (9.28, 12.60, 3.33), "rsr": (3.89, 3.89, None),
                   "rsnr": (0.01, 0.42, 0.41), "total": (13.18, 16.91, 3.74)},
    "2026-03-31": {"rrc": (9.17, 12.72, 3.55), "rsr": (4.19, 4.19, None),
                   "rsnr": (0.01, 0.56, 0.55), "total": (13.37, 17.46, 4.10)},
    "2026-06-30": {"rrc": (9.08, 12.80, 3.72), "rsr": (2.77, 2.77, None),
                   "rsnr": (0.01, 0.81, 0.80), "total": (11.87, 16.38, 4.52)},
}
INCREMENTO = {"rrc": 0.17, "rsr": None, "rsnr": 0.25, "total": 0.42}

fallos: "list[str]" = []
pruebas = 0


def igual(etiqueta, obtenido, esperado, tol):
    global pruebas
    pruebas += 1
    if obtenido is None and esperado is None:
        return
    if obtenido is None or esperado is None or abs(obtenido - esperado) > tol:
        fallos.append(f"{etiqueta}: se obtuvo {obtenido!r}, se esperaba {esperado!r}")


def main(argv):
    global pruebas
    balanza = Path(argv[0]) if argv else AQUI / "Balanza_062026.xlsx"
    actuarios = Path(argv[1]) if len(argv) > 1 else AQUI / "ResultadosQES.xlsb"
    for f in (balanza, actuarios):
        if not f.exists():
            print(f"No encuentro {f}. Uso: python verificar.py <balanza.xlsx> <actuarios.xlsb>")
            return 2

    tmp = Path(tempfile.mkdtemp(prefix="reservas-qes-"))

    # ---- 1. lectura de los dos archivos ---------------------------------
    h = Historico(tmp / "historico.json")
    for aviso in h.procesar(actuarios):
        print("·", aviso)
    for aviso in h.procesar(balanza):
        print("·", aviso)
    h.guardar()
    print()

    igual("número de cortes", len(h.periodos()), 6, 0)
    for p, (rrc_l, rrc_c, rsr, rsnr_l, rsnr_c) in ESPERADO.items():
        if p not in h.datos:
            fallos.append(f"falta el corte {p}")
            continue
        e = h.datos[p]
        igual(f"{p} rrc local",  e["local"].get("rrc"),  rrc_l,  0.005)
        igual(f"{p} rrc cnsf",   e["cnsf"].get("rrc"),   rrc_c,  0.005)
        igual(f"{p} rsr local",  e["local"].get("rsr"),  rsr,    0.005)
        igual(f"{p} rsr cnsf",   e["cnsf"].get("rsr"),   rsr,    0.005)
        igual(f"{p} rsnr local", e["local"].get("rsnr"), rsnr_l, 0.005)
        igual(f"{p} rsnr cnsf",  e["cnsf"].get("rsnr"),  rsnr_c, 0.005)

    # la balanza de junio debe cuadrar al centavo con la columna local de actuarios
    ref = h.datos["2026-06-30"].get("local_actuarios", {})
    for cid in ("rrc", "rsr", "rsnr"):
        igual(f"2026-06-30 balanza vs actuarios ({cid})",
              h.datos["2026-06-30"]["local"][cid], ref.get(cid), 0.005)
    if h.datos["2026-06-30"].get("aviso"):
        fallos.append("junio 2026 quedó marcado: " + h.datos["2026-06-30"]["aviso"])

    # ---- 2. la vista en millones ----------------------------------------
    ps = h.periodos()[-3:]
    igual("cortes en la vista", len(ps), 3, 0)
    for p, filas in VISTA.items():
        for cid, (loc, cnsf, dif) in filas.items():
            if cid == "total":
                igual(f"{p} total local", h.total(p, "local") / 1e6, loc, 0.005)
                igual(f"{p} total cnsf", h.total(p, "cnsf") / 1e6, cnsf, 0.005)
                igual(f"{p} total dif", h.diferencia(p) / 1e6, dif, 0.005)
            else:
                igual(f"{p} {cid} local", h.datos[p]["local"][cid] / 1e6, loc, 0.005)
                igual(f"{p} {cid} cnsf", h.datos[p]["cnsf"][cid] / 1e6, cnsf, 0.005)
                d = h.diferencia(p, cid) / 1e6
                igual(f"{p} {cid} dif", None if abs(d) < 0.005 else d, dif, 0.005)

    for cid, inc in INCREMENTO.items():
        v = ((h.diferencia(ps[-1]) - h.diferencia(ps[-2])) / 1e6 if cid == "total"
             else h.incremento(ps[-1], ps[-2], cid) / 1e6)
        igual(f"incremento {cid}", None if abs(v) < 0.005 else v, inc, 0.005)

    # ---- 3. indicadores --------------------------------------------------
    fx = BLOQUE["FX_DEFAULT"]
    d1 = h.diferencia("2026-06-30")
    d0 = h.diferencia("2026-03-31")
    igual("KPI diferencia total (MM USD)", d1 / 1e6, 4.52, 0.005)
    igual("KPI variación del trimestre (MM USD)", (d1 - d0) / 1e6, 0.42, 0.005)
    igual("KPI variación %", (d1 - d0) / d0 * 100, 10.3, 0.05)
    igual("KPI equivalente en MXN (MM)", d1 * fx / 1e6, 79.04, 0.01)

    # ---- 4. el HTML ------------------------------------------------------
    ruta = escribir_vista(h, periodos=ps, abrir=False)
    doc = ruta.read_text(encoding="utf-8")
    igual("nombre del archivo", 1 if ruta.name == "vista_reservas_2026-06-30.html" else 0, 1, 0)
    for cadena in ("Resultados reservas técnicas QES",   # el CSS lo pone en versalitas
                   "text-transform:uppercase",
                   "Comparativo de metodologías y evolución del diferencial",
                   "Incremento", "Total reservas",
                   "Esta reserva se calcula una vez al año",
                   "Mensajes clave", "Nota relevante", "<svg", "4.52", "10.3%", "0.42"):
        pruebas += 1
        if cadena not in doc:
            fallos.append(f"el HTML no trae «{cadena}»")
    for prohibido in ("http://", "https://", "<link", " src="):
        pruebas += 1
        if prohibido in doc:
            fallos.append(f"el HTML no es autocontenido: contiene «{prohibido}»")
    pruebas += 1
    if doc.count("—") < 3:      # el guion donde el valor es cero
        fallos.append("faltan los guiones donde el valor es cero")

    # ---- 5. el mes que entra: solo la balanza, el histórico se conserva ---
    h2 = Historico(tmp / "historico.json")
    igual("el histórico se releyó completo", len(h2.periodos()), 6, 0)
    antes = dict(h2.datos["2025-12-31"]["cnsf"])
    h2.procesar(balanza)          # como si fuera la balanza del mes siguiente
    igual("cortes tras recargar la balanza", len(h2.periodos()), 6, 0)
    igual("un corte anterior quedó intacto", h2.datos["2025-12-31"]["cnsf"]["rrc"],
          antes["rrc"], 0.005)

    # ---- 5 bis. la tablita de los actuarios, en las dos convenciones -----
    # Los importes pueden venir como número de Excel o como texto, y el texto
    # tanto «$8,331,317.86» como «$8.331.317,86». Leer el segundo como 8.33 no
    # truena: sólo mete un importe mil veces más chico en el reporte.
    to_num = BLOQUE["to_num"]
    NUMEROS = [
        ("$8.331.317,86", 8331317.86), ("$11.142.522,97", 11142522.97),
        ("$12.701,88", 12701.88), ("$325.011,19", 325011.19),
        ("$14.286.192,37 ", 14286192.37), ("8331317,86", 8331317.86),
        ("$8,331,317.86", 8331317.86), ("8,331,317.86", 8331317.86),
        ("8331317.86", 8331317.86), ("(1.234,50)", -1234.50),
        ("(1,234.50)", -1234.50), ("-8.331.317,86", -8331317.86),
        ("1,234", 1234.0), ("1.234", 1.234), ("12.70", 12.70),
        ("1.234.567", 1234567.0), ("1,234,567", 1234567.0), ("7461,06", 7461.06),
    ]
    for txt, esp in NUMEROS:
        igual(f"to_num({txt!r})", to_num(txt), esp, 0.005)
    for basura in ("", "abc", "1.2.3,4,5", "-", "$", "12/05/2026"):
        pruebas += 1
        if to_num(basura) is not None:
            fallos.append(f"to_num aceptó basura: {basura!r} -> {to_num(basura)!r}")

    # la tablita tal como la mandan los actuarios, armada aquí y leída de vuelta
    from openpyxl import Workbook
    def _eu(x):
        return "$" + f"{x:,.2f}".replace(",", "\x00").replace(".", ",").replace("\x00", ".") + " "
    BLOQUES = {
        "31 de Julio  de 2025":    [(8331317.86, 11142522.97), (2818658.21, 2818658.21),
                                    (12701.88, 325011.19)],
        "30 de Junio de 2026":     [(9084601.70, 12799633.26), (2773742.22, 2773742.22),
                                    (7461.06, 809405.38)],
    }
    NOMS = ["Reserva de Riesgos en Curso", "Reserva de Siniestros Reportados",
            "Reserva de Siniestros No reportados"]
    for etiqueta, fmt in (("número", None), ("texto europeo", _eu),
                          ("texto inglés", lambda x: f"${x:,.2f} ")):
        wb = Workbook(); ws = wb.active
        r = 2
        for titulo, filas in BLOQUES.items():
            ws.cell(r, 2, titulo)
            ws.cell(r + 1, 1, "Reserva")
            ws.cell(r + 1, 2, "QES\nMetodología local")
            ws.cell(r + 1, 3, "QES\nCNSF\nMetodo Estatutarío")
            for k, (nom, (loc, cn)) in enumerate(zip(NOMS, filas)):
                ws.cell(r + 2 + k, 1, nom)
                ws.cell(r + 2 + k, 2, fmt(loc) if fmt else loc)
                ws.cell(r + 2 + k, 3, fmt(cn) if fmt else cn)
            tl, tcn = sum(a for a, _ in filas), sum(b for _, b in filas)
            ws.cell(r + 5, 1, "Total Reservas")
            ws.cell(r + 5, 2, fmt(tl) if fmt else tl)
            ws.cell(r + 5, 3, fmt(tcn) if fmt else tcn)
            r += 8
        ruta_act = tmp / f"act_{etiqueta.split()[-1]}.xlsx"
        wb.save(ruta_act)

        lec = BLOQUE["leer_fuente"](ruta_act)
        igual(f"actuarios en {etiqueta}: se reconoce",
              1 if lec.fuente == "actuarios" else 0, 1, 0)
        igual(f"actuarios en {etiqueta}: cortes", len(lec.cortes), 2, 0)
        for per, filas in (("2025-07-31", BLOQUES["31 de Julio  de 2025"]),
                           ("2026-06-30", BLOQUES["30 de Junio de 2026"])):
            v = lec.periodos.get(per, {"local": {}, "cnsf": {}})
            igual(f"{etiqueta} · {per} conceptos", len(v["cnsf"]), 3, 0)
            for cid, (loc, cn) in zip(("rrc", "rsr", "rsnr"), filas):
                igual(f"{etiqueta} · {per} {cid} local", v["local"].get(cid), loc, 0.005)
                igual(f"{etiqueta} · {per} {cid} CNSF", v["cnsf"].get(cid), cn, 0.005)

    # la balanza manda para la columna local, y el desajuste se avisa
    import openpyxl as _op
    wb = _op.load_workbook(tmp / "act_europeo.xlsx"); ws = wb.active
    destino = max((c.row for f in ws.iter_rows() for c in f
                   if isinstance(c.value, str) and "Riesgos en Curso" in c.value))
    ws.cell(destino, 2).value = "$9.000.000,00 "
    wb.save(tmp / "act_desajustado.xlsx")
    for etq, orden in (("actuarios primero", [tmp / "act_desajustado.xlsx", balanza]),
                       ("balanza primero", [balanza, tmp / "act_desajustado.xlsx"])):
        hx = Historico(tmp / f"hx_{etq[0]}.json")
        for f in orden:
            hx.procesar(f)
        igual(f"{etq}: la columna local sale de la balanza",
              hx.datos["2026-06-30"]["local"]["rrc"], 9084601.70, 0.005)
        pruebas += 1
        aviso = hx.datos["2026-06-30"].get("aviso", "")
        if "no coincide" not in aviso or "9,000,000.00" not in aviso:
            fallos.append(f"{etq}: el check no avisó del desajuste (aviso: {aviso!r})")

    # ---- 6. el servidor: subir, catalogar y volver a armar --------------
    Repositorio, leer_fuente = BLOQUE["Repositorio"], BLOQUE["leer_fuente"]
    repo = Repositorio(url=f"sqlite:///{tmp/'audit.db'}")
    repo.conectar()
    igual("el esquema aún no existe", 1 if repo.existe_esquema() else 0, 0, 0)
    igual("el catálogo vacío no revienta", len(repo.snapshots()), 0, 0)
    creadas = repo.crear_esquema()
    for objeto in (BLOQUE["TABLA_CARGAS"], BLOQUE["TABLA_DETALLE"], BLOQUE["VISTA_SQL"]):
        pruebas += 1
        if objeto not in creadas:
            fallos.append(f"crear_esquema no creó «{objeto}»: creó {creadas}")
    igual("crear el esquema dos veces no duplica", len(repo.crear_esquema()), 0, 0)
    igual("el esquema ya existe", 1 if repo.existe_esquema() else 0, 1, 0)

    lec_act = leer_fuente(actuarios)
    id_act = repo.subir(lec_act, usuario="prueba")
    igual("cortes que trae el archivo de actuarios", len(lec_act.cortes), 6, 0)
    lec_bal = leer_fuente(balanza)
    id_bal = repo.subir(lec_bal, usuario="prueba")
    igual("la balanza se sube con su corte", len(lec_bal.cortes), 1, 0)
    igual("se reconoce el archivo ya subido", len(repo.ya_subido(lec_bal.sha256)), 1, 0)

    # una copia más del mismo mes: la anterior se conserva, no se pisa
    id_bis = repo.subir(lec_bal, usuario="otro", nota="segunda copia")
    pruebas += 1
    if id_bis == id_bal:
        fallos.append("la segunda copia pisó la primera en vez de quedar aparte")
    snaps = repo.snapshots()
    igual("copias de junio 2026 en el catálogo",
          len([s for s in snaps if s["periodo"] == "2026-06-30"]), 3, 0)
    igual("usuarios distintos en el catálogo",
          len({s["usuario"] for s in snaps}), 2, 0)

    # los importes vuelven del servidor idénticos a como entraron
    for cid in ("rrc", "rsr", "rsnr"):
        igual(f"ida y vuelta por la base ({cid})",
              repo.importes(id_bal, "2026-06-30")["local"][cid],
              h.datos["2026-06-30"]["local"][cid], 0.005)

    # las tres ranuras: diciembre · t-1 · t, con la copia elegida a mano
    hs = repo.historico([("2025-12-31", id_act), ("2026-03-31", id_act),
                         ("2026-06-30", id_bal)], ruta_json=tmp / "hs.json")
    igual("cortes armados desde el servidor", len(hs.periodos()), 3, 0)
    for per, filas in VISTA.items():
        igual(f"servidor · {per} total dif", hs.diferencia(per) / 1e6,
              filas["total"][2], 0.005)
    igual("servidor · variación %",
          (hs.diferencia("2026-06-30") - hs.diferencia("2026-03-31"))
          / hs.diferencia("2026-03-31") * 100, 10.3, 0.05)
    # el origen debe decir de qué carga exacta salió cada columna
    org = hs.datos["2026-06-30"]["origen"]
    for etq, texto, carga, archivo in (
            ("local", org.get("local", ""), id_bal, Path(balanza).name),
            ("cnsf", org.get("cnsf", ""), id_act, Path(actuarios).name)):
        pruebas += 1
        if f"#{carga}" not in texto or archivo not in texto:
            fallos.append(f"el origen de la columna {etq} no identifica su carga "
                          f"(esperaba #{carga} y {archivo}): {texto!r}")

    # ---- 7. la evolución mes con mes ------------------------------------
    escribir_evolucion = BLOQUE["escribir_evolucion"]
    ruta_ev = escribir_evolucion(h, destino=tmp / "evolucion.html", abrir=False)
    ev = ruta_ev.read_text(encoding="utf-8")
    for cadena in ("Evolución de la diferencia", "3.12", "4.52", "<svg", "MM USD",
                   "Las mismas cifras, en tabla", "Diferencia total"):
        pruebas += 1
        if cadena not in ev:
            fallos.append(f"la evolución no trae «{cadena}»")
    for prohibido in ("http://", "https://", "<link", " src="):
        pruebas += 1
        if prohibido in ev:
            fallos.append(f"la evolución no es autocontenida: contiene «{prohibido}»")
    # un globo por columna, y cada uno emparejado con la suya
    igual("columnas graficadas", ev.count('class="col c'), 6, 0)
    igual("globos del cursor", ev.count('class="tip t'), 6, 0)
    for i in range(6):
        pruebas += 1
        if f".c{i}:hover ~ .t{i}" not in ev:
            fallos.append(f"la columna {i} no enciende su globo")
    # los globos se pintan al final: si no, la columna siguiente los tapa
    pruebas += 1
    if ev.index('class="tip t0') < ev.rindex('class="col c'):
        fallos.append("los globos se dibujan antes que las columnas y quedan tapados")

    # ---- 7 bis. la base propia y la vista plana --------------------------
    from sqlalchemy import inspect as _inspect
    igual("la vista plana existe en la base",
          1 if BLOQUE["VISTA_SQL"] in _inspect(repo.engine).get_view_names() else 0, 1, 0)
    from sqlalchemy import text as _text
    with repo.engine.connect() as _c:
        fila = _c.execute(_text(
            "SELECT metodologia_local, metodo_estatutario, diferencia FROM "
            f"{BLOQUE['VISTA_SQL']} WHERE periodo='2026-06-30' AND concepto='rrc' "
            f"AND carga_id={id_act}")).one()
    igual("la vista trae la metodología local", float(fila[0]),
          h.datos["2026-06-30"]["local_actuarios"]["rrc"], 0.005)
    igual("la vista trae el método estatutario", float(fila[1]),
          h.datos["2026-06-30"]["cnsf"]["rrc"], 0.005)
    igual("la vista calcula la diferencia", float(fila[2]),
          h.diferencia("2026-06-30", "rrc"), 0.5)

    # el nombre de la base se interpola en el SQL: sólo identificadores simples
    for malo in ("Reservas QES", "1base", "x];DROP DATABASE y--", "a" * 130, "base-x"):
        pruebas += 1
        try:
            repo.crear_base(malo)
            fallos.append(f"aceptó un nombre de base inválido: {malo!r}")
        except ValueError:
            pass
        except Exception as err:
            fallos.append(f"el nombre {malo!r} falló con algo que no es ValueError: {err}")
    pruebas += 1
    if repo.crear_base("Reservas_QES") is not False:
        fallos.append("en SQLite crear_base debería no hacer nada y devolver False")

    # ---- 7 ter. la cascada encadena todos los cortes elegidos ------------
    # La dibuja el navegador cuando el lector mete o saca meses; la versión en
    # Python es el respaldo del <noscript> y la que se comprueba aquí.
    import json as _json
    import re as _r
    _re = _r
    svg_de = BLOQUE["_svg_cascada"]
    seis = h.periodos()
    for n in (2, 3, 4, 6):
        ps_ = seis[-n:]
        svg = svg_de(h, ps_, {p_: h.fx(p_) for p_ in ps_})
        mov = sum(1 for a, b in zip(ps_, ps_[1:]) for c in BLOQUE["CONCEPTOS"]
                  if abs((h.incremento(b, a, c.id) or 0)) >= 5000)
        igual(f"barras de la cascada con {n} cortes", svg.count("<rect "), n + mov, 0)
        for per in ps_:
            pruebas += 1
            esperado = f"{(h.diferencia(per) or 0) / 1e6:,.2f}"
            if f">{esperado}</text>" not in svg:
                fallos.append(f"la cascada no marca {esperado} para {etiqueta_corta(per)}")
        pruebas += 1
        ancho = int(_r.search(r'viewBox="0 0 (\d+)', svg).group(1))
        if f"min-width:{ancho}px" not in svg:
            fallos.append(f"la cascada de {n} cortes se dejaría comprimir")

    # el lienzo crece con los cortes: apretarlo encimaba las etiquetas del pie
    a2 = int(_r.search(r'viewBox="0 0 (\d+)', svg_de(h, seis[-2:], {p_: h.fx(p_) for p_ in seis[-2:]})).group(1))
    a6 = int(_r.search(r'viewBox="0 0 (\d+)', svg_de(h, seis, {p_: h.fx(p_) for p_ in seis})).group(1))
    pruebas += 1
    if a6 <= a2:
        fallos.append(f"el lienzo no crece con los cortes: {a2} con 2, {a6} con 6")

    # lo que baja va en rojo
    svg4 = svg_de(h, seis[-4:], {p_: h.fx(p_) for p_ in seis[-4:]})
    pruebas += 1
    if "−0.17" not in svg4 or BLOQUE["NEG"] not in svg4:
        fallos.append("la cascada no marca en rojo el movimiento que baja")

    # ---- 7 quater. las versiones master ----------------------------------
    igual("sin marcar nada no hay master", len(repo.maestros()), 0, 0)
    loc, cn = repo.master_sugerida("2026-06-30")
    igual("la master sugerida es la balanza más nueva", loc, id_bis, 0)
    igual("y su cnsf, los actuarios", cn, id_act, 0)
    repo.marcar_master("2026-06-30", id_bal, id_act, etiqueta="Cierre junio", usuario="prueba")
    repo.marcar_master("2026-03-31", id_act, id_act, etiqueta="Cierre marzo")
    igual("cortes master", len(repo.maestros()), 2, 0)
    hm = repo.historico_master(tmp / "hm.json")
    igual("el histórico master sólo trae los master", len(hm.periodos()), 2, 0)
    igual("y con sus cifras", hm.diferencia("2026-06-30") / 1e6, 4.52, 0.005)
    igual("la etiqueta viaja con el corte",
          1 if hm.datos["2026-06-30"].get("etiqueta") == "Cierre junio" else 0, 1, 0)
    # cambiar la master de un corte lo cambia en la vista, sin tocar las cargas
    repo.marcar_master("2026-06-30", id_bis, id_act, etiqueta="Cierre junio v2")
    igual("cambiar la master no duplica el corte", len(repo.maestros()), 2, 0)
    pruebas += 1
    if f"#{id_bis}" not in repo.historico_master(tmp / "hm2.json").datos["2026-06-30"]["origen"]["local"]:
        fallos.append("cambiar la master no cambió de qué carga sale la columna local")
    igual("quitar una master", 1 if repo.quitar_master("2026-03-31") else 0, 1, 0)
    igual("queda una", len(repo.maestros()), 1, 0)

    # ---- 7 quinquies. el HTML se rehace solo -----------------------------
    doc_i = construir_html(h, periodos=ps, disponibles=h.periodos())
    payload_i = _json.loads(_re.search(r"^const D = (\{.*\});$", doc_i, _re.M).group(1))
    igual("cortes que viajan dentro del HTML", len(payload_i["cortes"]), 6, 0)
    igual("cortes marcados al abrir", len(payload_i["activos"]), 3, 0)
    igual("casillas de mes", doc_i.count('<label class="mes'), 6, 0)
    # los textos que dependen de un par vienen escritos desde Python
    igual("pares precalculados", len(payload_i["pares"]), 15, 0)   # 6*5/2
    for a, b in (("2025-12-31", "2026-06-30"), ("2025-07-31", "2026-03-31")):
        pruebas += 1
        if f"{a}|{b}" not in payload_i["pares"]:
            fallos.append(f"falta el par precalculado {a}|{b}")
    pruebas += 1
    if "n2(" not in doc_i or "function cascada" not in doc_i:
        fallos.append("el HTML no trae el motor que lo rehace")

    # ---- 7 sexies. la letra, para quien batalla para ver -----------------
    css = BLOQUE["CSS"]
    tam = [float(x) for x in _re.findall(r"font-size:(\d+(?:\.\d+)?)px", css)]
    igual("el tamaño de letra más chico de la hoja", min(tam), 13.0, 0.5)
    pruebas += 1
    if min(tam) < 12.9:
        fallos.append(f"hay letra de {min(tam)}px: demasiado chica")
    for chico in _re.findall(r'font-size="(\d+(?:\.\d+)?)"', doc_i):
        pruebas += 1
        if float(chico) < 12:
            fallos.append(f"hay texto de {chico}px en un gráfico: demasiado chico")
    pruebas += 1
    if 'id="zoom"' not in doc_i:
        fallos.append("falta el control de tamaño de letra")

    # ---- 8. las dos monedas ---------------------------------------------
    TC = {"2025-12-31": 17.8410, "2026-03-31": 17.6220, "2026-06-30": 17.4986}
    for per, v in TC.items():
        h.poner_fx(per, v)
    igual("el TC de cierre se guarda", h.fx("2026-06-30"), 17.4986, 1e-9)
    igual("sin TC propio se usa el general", h.fx("2025-07-31", 19.0), 19.0, 1e-9)
    h.poner_fx("2026-06-30", None)
    igual("se puede borrar el TC de un corte", h.fx("2026-06-30", 19.0), 19.0, 1e-9)
    h.poner_fx("2026-06-30", 17.4986)

    ruta2 = escribir_vista(h, destino=tmp / "moneda.html", periodos=ps, abrir=False)
    doc2 = ruta2.read_text(encoding="utf-8")
    for cadena in ('id="m-usd"', 'id="m-mxn"', 'label for="m-mxn"',
                   'class="v-usd"', 'class="v-mxn"', "Dólares", "Pesos"):
        pruebas += 1
        if cadena not in doc2:
            fallos.append(f"al interruptor de moneda le falta «{cadena}»")
    for prohibido in ("http://", "https://", "<link", " src="):
        pruebas += 1
        if prohibido in doc2:
            fallos.append(f"el interruptor metió algo externo: «{prohibido}»")

    # cada corte convertido con SU tipo de cambio
    payload = _json.loads(_re.search(r"^const D = (\{.*\});$", doc2, _re.M).group(1))
    for per in ps:
        c = payload["cortes"][per]
        igual(f"datos embebidos · {per} tipo de cambio", c["fx"], TC[per], 1e-9)
        for cid in ("rrc", "rsr", "rsnr"):
            igual(f"datos embebidos · {per} {cid} local", c["local"][cid],
                  h.datos[per]["local"][cid], 0.005)
            igual(f"datos embebidos · {per} {cid} cnsf", c["cnsf"][cid],
                  h.datos[per]["cnsf"][cid], 0.005)

    # el KPI y el mensaje clave no pueden contradecirse
    v_usd = (h.diferencia("2026-06-30") - h.diferencia("2026-03-31")) / 1e6
    v_mxn = (h.diferencia("2026-06-30") * TC["2026-06-30"]
             - h.diferencia("2026-03-31") * TC["2026-03-31"]) / 1e6
    igual("variación en pesos con el TC de cada cierre", v_mxn, 6.88, 0.005)
    igual("variación en dólares", v_usd, 0.42, 0.005)
    msg_usd = h.mensajes("2026-06-30", "2026-03-31")[2]
    msg_mxn = h.mensajes("2026-06-30", "2026-03-31", fx=TC["2026-06-30"],
                         fx_previo=TC["2026-03-31"])[2]
    for etq, msg, cifra in (("dólares", msg_usd, f"USD {v_usd:,.2f} MM"),
                            ("pesos", msg_mxn, f"MXN {v_mxn:,.2f} MM")):
        pruebas += 1
        if cifra not in msg:
            fallos.append(f"el mensaje en {etq} no dice «{cifra}»: {msg[-130:]!r}")
    pruebas += 1
    if "+9.5%" not in msg_mxn:
        fallos.append(f"el porcentaje en pesos debería ser +9.5%: {msg_mxn[-90:]!r}")

    # los gráficos también traen las dos monedas
    pruebas += 1
    if 'class="v-mxn" x=' not in doc2:
        fallos.append("la cascada no trae las etiquetas en pesos")
    ev2 = escribir_evolucion(h, destino=tmp / "ev2.html", abrir=False).read_text(encoding="utf-8")
    for cadena in ('id="m-mxn"', "MM MXN", 'class="v-mxn"'):
        pruebas += 1
        if cadena not in ev2:
            fallos.append(f"la evolución no trae «{cadena}»")

    print(f"{pruebas} comprobaciones · {len(fallos)} fallo(s)")
    if fallos:
        for f in fallos:
            print("  ✗", f)
        return 1
    print("Todo cuadra. HTML de muestra:", ruta)
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
